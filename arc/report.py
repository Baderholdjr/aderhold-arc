"""Client reports: scorecard JSON, a self-contained HTML readout and an Excel scorecard.

Reports contain aggregates only (groups of 5+). They are written to an output folder
that is git-ignored; the nightly workflow uploads them as private workflow artifacts.
"""
from __future__ import annotations

import html
import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .benchmark import percentile_rank
from .roi import RoiResult
from .scoring import CampaignScore

SEV_COLOR = {"Critical": "#c0392b", "Significant": "#d68910", "Moderate": "#b7950b", "Healthy": "#1e8449"}


def _f(v, digits=0, pct=False, money=False):
    if v is None:
        return "n<5" if not money else "-"
    if pct:
        return f"{v * 100:.{digits}f}%"
    if money:
        return f"${v:,.0f}"
    return f"{v:.{digits}f}"


def _cell_bg(v):
    if v is None:
        return "#f2f2f2"
    # red (20) -> amber (60) -> green (90)
    t = max(0.0, min(1.0, (v - 20) / 70))
    if t < 4 / 7:
        u = t / (4 / 7)
        r, g, b = 248 + (255 - 248) * u, 105 + (235 - 105) * u, 107 + (132 - 107) * u
    else:
        u = (t - 4 / 7) / (3 / 7)
        r, g, b = 255 + (99 - 255) * u, 235 + (190 - 235) * u, 132 + (123 - 132) * u
    return f"rgb({int(r)},{int(g)},{int(b)})"


def to_json(score: CampaignScore, roi: RoiResult | None, path: Path) -> None:
    payload = {"generated": date.today().isoformat(), "score": score.to_dict(),
               "roi": roi.to_dict() if roi else None}
    path.write_text(json.dumps(payload, indent=2))


def to_html(score: CampaignScore, roi: RoiResult | None, client: str, bench: dict | None, path: Path) -> None:
    e = html.escape
    rows = []
    for d in score.dimensions:
        pos = percentile_rank(d.index_all, d.code, bench) or ""
        sev = d.severity or ""
        rows.append(
            f"<tr><td>{e(d.code)}</td><td>{e(d.name)}</td><td>{e(d.pillar)}</td>"
            f"<td style='background:{_cell_bg(d.index_all)}'>{_f(d.index_all)}</td>"
            f"<td style='background:{_cell_bg(d.index_ic)}'>{_f(d.index_ic)}</td>"
            f"<td style='background:{_cell_bg(d.index_leaders)}'>{_f(d.index_leaders)}</td>"
            f"<td>{_f(d.perception_gap) if d.perception_gap is not None else ''}</td>"
            f"<td style='color:{SEV_COLOR.get(sev, '#333')};font-weight:600'>{e(sev)}</td>"
            f"<td>{e(d.flag or '')}</td><td>{e(pos)}</td></tr>")
    pillars = "".join(
        f"<div class='pill'><div class='pn'>{e(p.pillar)}</div><div class='pv'>{_f(p.index_all)}</div>"
        f"<div class='pl'>Level {p.level or '-'}: {e(p.level_name or '')}</div></div>" for p in score.pillars)
    gaps = "".join(f"<li><b>{e(g['code'])} {e(g['name'])}</b>: {_f(g['index'])} ({e(g['severity'] or '')})</li>"
                   for g in score.top_gaps)
    b = score.behavioral or {}
    beh = [("Active adoption (weekly or more)", _f(b.get("active_adoption"), pct=True)),
           ("Depth of use", _f(b.get("depth_of_use"), pct=True)),
           ("Share of saved time spent re-checking", _f(b.get("verification_share"), pct=True)),
           ("Saw an AI error reach output (90 days)", _f(b.get("error_seen_90d"), pct=True)),
           ("Average AI training hours (12 months)", _f(b.get("training_hours_avg"), 1)),
           ("Under 5 hours of training", _f(b.get("share_under_5_training_hours"), pct=True)),
           ("Own-time learning hours (3 months)", _f(b.get("own_time_hours_avg"), 1)),
           ("Used unapproved AI tools (30 days)", _f(b.get("shadow_ai_30d"), pct=True))]
    beh_html = "".join(f"<tr><td>{e(k)}</td><td>{v}</td></tr>" for k, v in beh)

    roi_html = ""
    if roi:
        mx = (roi.achievable_value or 1) / 0.8  # bars use at most 80% of the track; labels sit to their right
        bars = ""
        for w in roi.waterfall:
            if w["lost"] is None:
                left, width, cls, label = 0.0, w["running"] / mx * 100, "total", _f(w["running"], money=True)
            else:
                left, width, cls, label = w["running"] / mx * 100, w["lost"] / mx * 100, "lost", "-" + _f(w["lost"], money=True)
            bars += (f"<div class='wf'><span class='wl'>{e(w['step'])}</span><span class='track'>"
                     f"<span class='bar {cls}' style='left:{left:.1f}%;width:{width:.1f}%'></span>"
                     f"<span class='wv' style='left:calc({left + width:.1f}% + 6px)'>{label}</span></span></div>")
        roi_html = f"""
<h2>Value leakage</h2>
<p>People-side gaps let through <b>{_f(roi.share_realized, 1, pct=True)}</b> of achievable AI value
({_f(roi.current_value, money=True)} of {_f(roi.achievable_value, money=True)} a year).</p>
<div class='wfbox'>{bars}</div>
<table class='kv'>
<tr><td>Recoverable value per year</td><td><b>{_f(roi.recoverable_value, money=True)}</b></td></tr>
<tr><td>Risk-adjusted (after confidence haircut)</td><td>{_f(roi.risk_adjusted_value, money=True)}</td></tr>
<tr><td>Engagement fee</td><td>{_f(roi.engagement_fee, money=True)}</td></tr>
<tr><td>Net first-year return (unadjusted / risk-adjusted)</td><td>{_f(roi.net_return_x, 1)}x / {_f(roi.net_return_risk_adjusted_x, 1)}x</td></tr>
<tr><td>Payback</td><td>{_f(roi.payback_months, 1)} months</td></tr>
</table>
<p class='note'>Estimates use the client's inputs and survey-derived factors; see the ROI inputs in the campaign file.</p>"""

    supp = "<p class='warn'>Fewer than 5 responses: results suppressed to protect anonymity.</p>" if score.suppressed else ""
    doc = f"""<!doctype html><html lang='en'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>
<title>Aderhold A.R.C. Scorecard: {e(client)}</title>
<style>
body{{font-family:Arial,Helvetica,sans-serif;color:#1d2733;max-width:1080px;margin:32px auto;padding:0 16px;line-height:1.45}}
h1{{font-size:26px;margin:0}} h2{{font-size:19px;margin-top:32px;border-bottom:2px solid #1f3a5f;padding-bottom:4px}}
.sub{{color:#5b6b7c}} table{{border-collapse:collapse;width:100%;font-size:13px}}
th{{background:#1f3a5f;color:#fff;text-align:left;padding:6px}} td{{padding:5px 6px;border-bottom:1px solid #e3e7ec}}
.pills{{display:flex;gap:12px;flex-wrap:wrap}} .pill{{flex:1;min-width:180px;border:1px solid #cfd7e0;border-radius:8px;padding:12px}}
.pn{{font-weight:700}} .pv{{font-size:30px;font-weight:700;color:#1f3a5f}} .pl{{color:#5b6b7c;font-size:13px}}
.wf{{display:flex;align-items:center;gap:8px;font-size:13px;margin:4px 0}} .wl{{width:170px;flex:none}}
.track{{position:relative;flex:1;height:18px}} .wfbox .bar{{position:absolute;top:1px;height:16px}} .wv{{position:absolute;top:0;white-space:nowrap}}
.total{{background:#9aa5b1}} .lost{{background:#e6a19a}}
.kv td:first-child{{width:60%}} .note{{color:#5b6b7c;font-size:12px}} .warn{{color:#c0392b;font-weight:700}}
</style></head><body>
<h1>Aderhold A.R.C. Scorecard</h1>
<div class='sub'>{e(client)} &middot; campaign {e(score.campaign)} &middot; {score.n} respondents ({score.n_ic} individual contributors, {score.n_leaders} managers and leaders) &middot; instrument v{e(score.instrument_version)} &middot; generated {date.today().isoformat()}</div>
{supp}
<h2>Maturity by pillar</h2><div class='pills'>{pillars}</div>
<p>Overall A.R.C. index: <b>{_f(score.overall_index)}</b> / 100</p>
<h2>Top gaps</h2><ol>{gaps}</ol>
<h2>Gap dimensions</h2>
<table><tr><th>Code</th><th>Dimension</th><th>Pillar</th><th>All</th><th>IC</th><th>Leaders</th><th>Gap (pts)</th><th>Severity</th><th>Flag</th><th>Benchmark</th></tr>
{''.join(rows)}</table>
<p class='note'>Index 0-100. Severity: Critical &lt;40, Significant 40-59, Moderate 60-74, Healthy &ge;75. A perception gap means leaders score the dimension 15+ points above individual contributors. Groups under 5 are suppressed.</p>
<h2>Behavioral signals</h2><table class='kv'>{beh_html}</table>
{roi_html}
<p class='note'>Aderhold A.R.C.&trade; research-based diagnostic. Confidential: prepared for {e(client)}.</p>
</body></html>"""
    path.write_text(doc)


def to_xlsx(score: CampaignScore, roi: RoiResult | None, client: str, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Scorecard"
    hdr_font, hdr_fill = Font(name="Arial", bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3A5F")
    ws.append([f"Aderhold A.R.C. Scorecard: {client} ({score.campaign})"])
    ws["A1"].font = Font(name="Arial", bold=True, size=14)
    ws.append([f"n = {score.n} (IC {score.n_ic}, leaders {score.n_leaders}); instrument v{score.instrument_version}"])
    ws.append([])
    ws.append(["Code", "Dimension", "Pillar", "Index (all)", "Index (IC)", "Index (leaders)", "Perception gap", "Severity", "Flag"])
    for c in ws[4]:
        c.font, c.fill = hdr_font, hdr_fill
    for d in score.dimensions:
        ws.append([d.code, d.name, d.pillar, _r(d.index_all), _r(d.index_ic), _r(d.index_leaders),
                   _r(d.perception_gap), d.severity, d.flag])
    ws.append([])
    ws.append(["Pillar", "Index (all)", "Index (IC)", "Index (leaders)", "Level", "Level name"])
    for c in ws[ws.max_row]:
        c.font, c.fill = hdr_font, hdr_fill
    for p in score.pillars:
        ws.append([p.pillar, _r(p.index_all), _r(p.index_ic), _r(p.index_leaders), p.level, p.level_name])
    ws.append(["Overall", _r(score.overall_index)])
    for col, w in zip("ABCDEFGHI", [10, 36, 14, 12, 12, 14, 14, 12, 16]):
        ws.column_dimensions[col].width = w
    if roi:
        r = wb.create_sheet("Value leakage")
        for k, v in [("Achievable value", roi.achievable_value), ("Current realized value", roi.current_value),
                     ("Target value", roi.target_value), ("Share realized", roi.share_realized),
                     ("Recoverable value per year", roi.recoverable_value), ("Risk-adjusted value", roi.risk_adjusted_value),
                     ("Engagement fee", roi.engagement_fee), ("Net return (x)", roi.net_return_x),
                     ("Net return risk-adjusted (x)", roi.net_return_risk_adjusted_x), ("Payback (months)", roi.payback_months)]:
            r.append([k, v])
        r.column_dimensions["A"].width = 34
        r.column_dimensions["B"].width = 16
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.font.name != "Arial":
                    c.font = Font(name="Arial", bold=c.font.bold, size=c.font.size, color=c.font.color)
    wb.save(path)


def _r(v):
    return None if v is None else round(v, 1)
